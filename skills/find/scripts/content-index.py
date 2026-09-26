#!/usr/bin/env python3
"""Build and query a read-only local document content index for the find skill."""

from __future__ import annotations

import argparse
import datetime as dt
import fnmatch
import html
import ipaddress
import json
import logging
import os
from pathlib import Path
import re
import sqlite3
import subprocess
import sys
import urllib.parse
import urllib.request
import zipfile
import xml.etree.ElementTree as ET


SKILL_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_EXTENSIONS = "pdf,ppt,pptx,doc,docx,hwp,hwpx,xls,xlsx"
DEFAULT_ENDPOINT = "http://127.0.0.1:8080/"
DEFAULT_DB = Path(os.environ.get("LOCALAPPDATA", str(Path.home() / ".local" / "share"))) / "find-skill" / "content-index.db"


def first_config_line(name: str) -> str | None:
    path = SKILL_ROOT / "config" / name
    if not path.is_file():
        return None
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        value = line.strip().strip('"')
        if value and not value.startswith("#"):
            return value
    return None


def resolve_db(explicit: str | None) -> Path:
    value = explicit or os.environ.get("EVERYTHING_CONTENT_DB") or first_config_line("index-path.txt") or DEFAULT_DB
    return Path(os.path.expandvars(value)).expanduser().resolve()


def resolve_endpoint(explicit: str | None) -> str:
    value = explicit or os.environ.get("EVERYTHING_HTTP_URL") or first_config_line("http-url.txt") or DEFAULT_ENDPOINT
    value = value.strip().strip('"')
    parsed = urllib.parse.urlsplit(value)
    try:
        loopback = parsed.hostname == "localhost" or ipaddress.ip_address(parsed.hostname or "").is_loopback
        parsed.port  # Reject malformed port numbers before making a request.
    except ValueError:
        loopback = False
    if parsed.scheme not in {"http", "https"} or not loopback or parsed.username is not None or parsed.password is not None or parsed.query or parsed.fragment:
        raise ValueError("Use a loopback HTTP(S) endpoint without credentials, query, or fragment")
    return value.rstrip("/") + "/"


class NoEverythingRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        raise RuntimeError("Everything HTTP redirects are not allowed; check the endpoint")


def open_everything_url(url: str):
    return urllib.request.build_opener(NoEverythingRedirect).open(url, timeout=20)


def load_settings() -> dict:
    path = SKILL_ROOT / "config" / "index-config.json"
    if path.is_file():
        return json.loads(path.read_text(encoding="utf-8-sig"))
    return {}


def resolve_drives(explicit: str | None) -> list[str]:
    config = SKILL_ROOT / "config" / "search-config.json"
    configured = json.loads(config.read_text(encoding="utf-8-sig")).get("drives", []) if config.exists() else []
    drives = explicit.split(",") if explicit is not None else configured
    if not drives:
        drives = [os.environ.get("SystemDrive", "C:").rstrip(":\\")]
    if not isinstance(drives, list) or any(not isinstance(d, str) or not re.fullmatch(r"[A-Za-z]", d) for d in drives):
        raise ValueError("drives must be single letters, e.g. C,D")
    return sorted({d.upper() for d in drives})


def connect(db_path: Path, writable: bool = True) -> sqlite3.Connection:
    if not writable:
        uri = db_path.resolve().as_uri()
        wal_path = Path(str(db_path) + "-wal")
        if wal_path.exists():
            db = sqlite3.connect(f"{uri}?mode=ro", uri=True)
        else:
            db = sqlite3.connect(f"{uri}?mode=ro&immutable=1", uri=True)
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA query_only=ON")
        return db
    db_path.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(db_path)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA journal_mode=WAL")
    db.execute("PRAGMA foreign_keys=ON")
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS files (
            id INTEGER PRIMARY KEY,
            path TEXT NOT NULL UNIQUE,
            name TEXT NOT NULL,
            extension TEXT NOT NULL,
            size_bytes INTEGER NOT NULL,
            modified_ns INTEGER NOT NULL,
            indexed_utc TEXT NOT NULL,
            extractor TEXT,
            extraction_error TEXT,
            content TEXT NOT NULL DEFAULT ''
        );
        CREATE VIRTUAL TABLE IF NOT EXISTS files_fts USING fts5(
            name,
            path,
            content,
            content='files',
            content_rowid='id',
            tokenize='unicode61 remove_diacritics 2'
        );
        CREATE TRIGGER IF NOT EXISTS files_ai AFTER INSERT ON files BEGIN
          INSERT INTO files_fts(rowid, name, path, content)
          VALUES (new.id, new.name, new.path, new.content);
        END;
        CREATE TRIGGER IF NOT EXISTS files_ad AFTER DELETE ON files BEGIN
          INSERT INTO files_fts(files_fts, rowid, name, path, content)
          VALUES ('delete', old.id, old.name, old.path, old.content);
        END;
        CREATE TRIGGER IF NOT EXISTS files_au AFTER UPDATE ON files BEGIN
          INSERT INTO files_fts(files_fts, rowid, name, path, content)
          VALUES ('delete', old.id, old.name, old.path, old.content);
          INSERT INTO files_fts(rowid, name, path, content)
          VALUES (new.id, new.name, new.path, new.content);
        END;
        CREATE TABLE IF NOT EXISTS metadata (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        );
        """
    )
    return db


def decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def normalize_text(value: str, max_chars: int) -> str:
    value = value.replace("\x00", " ")
    value = re.sub(r"[\t\f\v ]+", " ", value)
    value = re.sub(r"\r\n?", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()[:max_chars]


def extract_zip_xml(path: Path, prefixes: tuple[str, ...]) -> str:
    parts: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            normalized = name.replace("\\", "/")
            if not normalized.lower().endswith(".xml") or not normalized.startswith(prefixes):
                continue
            try:
                root = ET.fromstring(archive.read(name))
                parts.extend(text.strip() for text in root.itertext() if text.strip())
            except ET.ParseError:
                continue
    return "\n".join(parts)


def extract_docx(path: Path) -> str:
    try:
        from docx import Document
        document = Document(path)
        parts = [p.text for p in document.paragraphs if p.text]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except ImportError:
        return extract_zip_xml(path, ("word/",))


def extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
        deck = Presentation(path)
        parts: list[str] = []
        for number, slide in enumerate(deck.slides, 1):
            slide_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_parts.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        slide_parts.append("\t".join(cell.text for cell in row.cells))
            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_parts.append(notes)
            except Exception:
                pass
            if slide_parts:
                parts.append(f"[Slide {number}]\n" + "\n".join(slide_parts))
        return "\n".join(parts)
    except ImportError:
        return extract_zip_xml(path, ("ppt/slides/", "ppt/notesSlides/"))


def extract_xlsx(path: Path, max_cells: int) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is unavailable") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    count = 0
    try:
        for sheet in workbook.worksheets:
            parts.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                count += len(values)
                if values:
                    parts.append("\t".join(values))
                if count >= max_cells:
                    parts.append("[cell limit reached]")
                    return "\n".join(parts)
    finally:
        workbook.close()
    return "\n".join(parts)


def extract_pdf(path: Path, max_pages: int) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is unavailable") from exc
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise RuntimeError("encrypted PDF") from exc
    parts: list[str] = []
    for page_number, page in enumerate(reader.pages[:max_pages], 1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"[Page {page_number}]\n{text}")
    return "\n".join(parts)


def resolve_tika() -> Path | None:
    value = os.environ.get("TIKA_APP_PATH") or first_config_line("tika-path.txt")
    if value and Path(os.path.expandvars(value)).is_file():
        return Path(os.path.expandvars(value))
    return None


def extract_hwp(path: Path, timeout: int) -> str:
    tika = resolve_tika()
    if not tika:
        raise RuntimeError("Apache Tika app not configured; set TIKA_APP_PATH or config/tika-path.txt")
    result = subprocess.run(
        ["java", "-jar", str(tika), "-t", str(path)],
        capture_output=True,
        timeout=timeout,
        check=False,
    )
    if result.returncode != 0:
        message = decode_text(result.stderr).strip() or f"Tika exit code {result.returncode}"
        raise RuntimeError(message[-1000:])
    return decode_text(result.stdout)


def extract_content(path: Path, settings: dict) -> tuple[str, str]:
    extension = path.suffix.lower().lstrip(".")
    max_chars = int(settings.get("max_text_chars", 2_000_000))
    if extension in {"txt", "md", "csv", "log", "json", "xml"}:
        text, extractor = decode_text(path.read_bytes()), "plain-text"
    elif extension in {"html", "htm"}:
        raw = decode_text(path.read_bytes())
        text = html.unescape(re.sub(r"<[^>]+>", " ", raw))
        extractor = "html"
    elif extension == "docx":
        text, extractor = extract_docx(path), "python-docx"
    elif extension in {"pptx", "pptm"}:
        text, extractor = extract_pptx(path), "python-pptx"
    elif extension in {"xlsx", "xlsm"}:
        text = extract_xlsx(path, int(settings.get("max_spreadsheet_cells", 200_000)))
        extractor = "openpyxl"
    elif extension == "pdf":
        text = extract_pdf(path, int(settings.get("max_pdf_pages", 1000)))
        extractor = "pypdf"
    elif extension == "hwpx":
        text, extractor = extract_zip_xml(path, ("Contents/", "Preview/")), "hwpx-xml"
    elif extension in {"hwp", "doc", "ppt", "xls"}:
        text, extractor = extract_hwp(path, int(settings.get("extract_timeout_seconds", 120))), "apache-tika"
    else:
        raise RuntimeError(f"unsupported extension: {extension}")
    return normalize_text(text, max_chars), extractor


def everything_results(endpoint: str, query: str, page_size: int):
    endpoint = resolve_endpoint(endpoint)
    offset = 0
    total = None
    observed_paths: set[str] = set()
    while total is None or offset < total:
        params = {
            "search": query,
            "json": "1",
            "path": "1",
            "path_column": "1",
            "size_column": "1",
            "date_modified_column": "1",
            "offset": str(offset),
            "count": str(page_size),
            "sort": "path",
            "ascending": "1",
        }
        url = endpoint + "?" + urllib.parse.urlencode(params)
        with open_everything_url(url) as response:
            payload = json.load(response)
        if not isinstance(payload, dict) or not isinstance(payload.get("results"), list):
            raise RuntimeError("Invalid Everything JSON response; pruning was not performed")
        raw_total = payload.get("totalResults")
        if isinstance(raw_total, bool) or not isinstance(raw_total, (int, str)) or not re.fullmatch(r"\d+", str(raw_total)):
            raise RuntimeError("Invalid Everything totalResults; pruning was not performed")
        page_total = int(raw_total)
        if total is not None and page_total != total:
            raise RuntimeError("Everything results changed during pagination; retry without pruning")
        total = page_total
        items = payload["results"]
        if offset + len(items) > total:
            raise RuntimeError("Inconsistent Everything result count; pruning was not performed")
        for item in items:
            if (not isinstance(item, dict) or not isinstance(item.get("path"), str) or not Path(item["path"]).is_absolute()
                    or not isinstance(item.get("name"), str) or not item["name"] or item["name"] in {".", ".."}
                    or "/" in item["name"] or "\\" in item["name"]):
                raise RuntimeError("Invalid Everything file path; pruning was not performed")
            path = Path(item["path"]) / item["name"]
            key = str(path).casefold()
            if key in observed_paths:
                raise RuntimeError("Duplicate Everything results during pagination; retry without pruning")
            observed_paths.add(key)
            yield path, total
        if not items and offset < total:
            raise RuntimeError("Everything returned an incomplete result set; pruning was not performed")
        offset += len(items)


def is_excluded(path: Path, patterns: list[str]) -> bool:
    normalized = str(path).replace("\\", "/").lower()
    return any(fnmatch.fnmatch(normalized, pattern.replace("\\", "/").lower()) for pattern in patterns)


def cmd_update(args: argparse.Namespace) -> int:
    settings = load_settings()
    extensions = [x.strip().lstrip(".").lower() for x in args.extensions.split(",") if x.strip()]
    if args.query and args.drives:
        raise ValueError("Use either --query or --drives, not both")
    if args.prune and args.query:
        raise ValueError("--prune cannot be used with a custom --query; use a dedicated database for scoped indexing")
    drives = resolve_drives(args.drives)
    if args.prune:
        missing_drives = [drive for drive in drives if not Path(drive + ":\\").is_dir()]
        if missing_drives:
            raise ValueError("Cannot prune while configured drives are inaccessible: " + ", ".join(missing_drives))
    drive_query = "|".join(d.lower() + ":" for d in drives)
    query = args.query or f"<{drive_query}> file: ext:{';'.join(sorted(set(extensions)))}"
    db_path = resolve_db(args.db)
    db = connect(db_path)
    max_bytes = int(args.max_mb * 1024 * 1024)
    patterns = list(settings.get("exclude_patterns", []))
    scope = json.dumps({"query": query, "excludes": sorted(patterns), "max_bytes": max_bytes}, sort_keys=True)
    counters = {"discovered": 0, "indexed": 0, "unchanged": 0, "skipped": 0, "errors": 0, "removed": 0}
    seen: set[str] = set()
    now = dt.datetime.now(dt.timezone.utc).isoformat()
    try:
        previous = dict(db.execute("SELECT key,value FROM metadata"))
        has_rows = db.execute("SELECT EXISTS(SELECT 1 FROM files)").fetchone()[0]
        mixed_scope = previous.get("mixed_scope") == "1" or bool(has_rows and previous.get("index_scope") != scope)
        if args.prune and mixed_scope:
            raise ValueError("--prune requires a database indexed with only this same scope; use a separate database for changed scopes")
        db.execute("INSERT OR REPLACE INTO metadata(key,value) VALUES('index_scope',?)", (scope,))
        db.execute("INSERT OR REPLACE INTO metadata(key,value) VALUES('mixed_scope',?)", ("1" if mixed_scope else "0",))
        for path, total in everything_results(resolve_endpoint(args.endpoint), query, args.page_size):
            counters["discovered"] += 1
            if counters["discovered"] == 1 or counters["discovered"] % args.progress_every == 0:
                print(f"[find-index] {counters['discovered']}/{total}: {path}", file=sys.stderr, flush=True)
            path_string = str(path)
            seen.add(path_string.casefold())
            try:
                if is_excluded(path, patterns) or not path.is_file():
                    counters["skipped"] += 1
                    continue
                stat = path.stat()
                if stat.st_size > max_bytes:
                    counters["skipped"] += 1
                    continue
                existing = db.execute("SELECT size_bytes, modified_ns, extraction_error FROM files WHERE path=?", (path_string,)).fetchone()
                if existing and not existing["extraction_error"] and existing["size_bytes"] == stat.st_size and existing["modified_ns"] == stat.st_mtime_ns:
                    counters["unchanged"] += 1
                    continue
                error = None
                extractor = None
                content = ""
                try:
                    content, extractor = extract_content(path, settings)
                except Exception as exc:
                    error = str(exc)[:2000]
                    counters["errors"] += 1
                db.execute(
                    """
                    INSERT INTO files(path,name,extension,size_bytes,modified_ns,indexed_utc,extractor,extraction_error,content)
                    VALUES(?,?,?,?,?,?,?,?,?)
                    ON CONFLICT(path) DO UPDATE SET
                      name=excluded.name, extension=excluded.extension, size_bytes=excluded.size_bytes,
                      modified_ns=excluded.modified_ns, indexed_utc=excluded.indexed_utc,
                      extractor=excluded.extractor, extraction_error=excluded.extraction_error, content=excluded.content
                    """,
                    (path_string, path.name, path.suffix.lower().lstrip("."), stat.st_size, stat.st_mtime_ns,
                     now, extractor, error, content),
                )
                counters["indexed"] += 1
                if counters["indexed"] % 20 == 0:
                    db.commit()
            except (OSError, PermissionError):
                counters["errors"] += 1
        if args.prune:
            rows = db.execute("SELECT id,path FROM files").fetchall()
            for row in rows:
                if row["path"].casefold() not in seen:
                    db.execute("DELETE FROM files WHERE id=?", (row["id"],))
                    counters["removed"] += 1
        db.execute("INSERT OR REPLACE INTO metadata(key,value) VALUES('last_update_utc',?)", (now,))
        db.execute("INSERT OR REPLACE INTO metadata(key,value) VALUES('last_query',?)", (query,))
        db.commit()
    finally:
        db.close()
    print(json.dumps({"database": str(db_path), "query": query, **counters}, ensure_ascii=False, indent=2))
    return 0


def fts_expression(query: str) -> str:
    terms = re.findall(r'"([^"]+)"|([^\s]+)', query)
    cleaned: list[str] = []
    for phrase, word in terms:
        value = (phrase or word).strip().replace('"', '""')
        if value:
            cleaned.append(f'"{value}"')
    if not cleaned:
        raise ValueError("search query is empty")
    return " AND ".join(cleaned)


def cmd_search(args: argparse.Namespace) -> int:
    db_path = resolve_db(args.db)
    if not db_path.is_file():
        raise FileNotFoundError(f"content index does not exist: {db_path}")
    db = connect(db_path, writable=False)
    filters = ["files_fts MATCH ?"]
    values: list[object] = [fts_expression(args.query)]
    if args.extensions:
        extensions = [x.strip().lower().lstrip(".") for x in args.extensions.split(",") if x.strip()]
        filters.append("files.extension IN (" + ",".join("?" for _ in extensions) + ")")
        values.extend(extensions)
    if args.path:
        def escape_like(text):
            return text.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        if Path(args.path).is_absolute():
            path = str(Path(args.path)).rstrip("\\/")
            filters.append("(files.path = ? COLLATE NOCASE OR files.path LIKE ? ESCAPE '\\')")
            values.extend([path, escape_like(path + os.sep) + "%"])
        else:
            filters.append("files.path LIKE ? ESCAPE '\\'")
            values.append("%" + escape_like(args.path) + "%")
    total = db.execute(
        f"SELECT COUNT(*) FROM files_fts JOIN files ON files.id=files_fts.rowid WHERE {' AND '.join(filters)}", values
    ).fetchone()[0]
    values.append(args.limit)
    rows = db.execute(
        f"""
        SELECT files.path, files.name, files.extension, files.size_bytes, files.modified_ns,
               files.extractor, files.extraction_error,
               snippet(files_fts, 2, '[', ']', ' … ', 32) AS excerpt,
               bm25(files_fts, 2.0, 0.5, 1.0) AS score
        FROM files_fts JOIN files ON files.id=files_fts.rowid
        WHERE {' AND '.join(filters)}
        ORDER BY score, files.modified_ns DESC
        LIMIT ?
        """,
        values,
    ).fetchall()
    results = []
    for row in rows:
        item = dict(row)
        item["modified_utc"] = dt.datetime.fromtimestamp(item.pop("modified_ns") / 1_000_000_000, dt.timezone.utc).isoformat()
        results.append(item)
    db.close()
    print(json.dumps({"database": str(db_path), "query": args.query, "total_results": total, "count": len(results), "results": results}, ensure_ascii=False, indent=2))
    return 0


def cmd_status(args: argparse.Namespace) -> int:
    db_path = resolve_db(args.db)
    if not db_path.is_file():
        print(json.dumps({"database": str(db_path), "exists": False}, ensure_ascii=False, indent=2))
        return 0
    db = connect(db_path, writable=False)
    totals = db.execute(
        "SELECT COUNT(*) files, SUM(CASE WHEN content<>'' THEN 1 ELSE 0 END) searchable, "
        "SUM(CASE WHEN extraction_error IS NOT NULL THEN 1 ELSE 0 END) errors, COALESCE(SUM(LENGTH(content)),0) text_chars FROM files"
    ).fetchone()
    by_extension = [dict(row) for row in db.execute("SELECT extension,COUNT(*) count FROM files GROUP BY extension ORDER BY count DESC")]
    metadata = {row["key"]: row["value"] for row in db.execute("SELECT key,value FROM metadata")}
    db.close()
    print(json.dumps({"database": str(db_path), "exists": True, **dict(totals), "by_extension": by_extension, **metadata}, ensure_ascii=False, indent=2))
    return 0


def positive_int(value: str) -> int:
    number = int(value)
    if number < 1:
        raise argparse.ArgumentTypeError("must be a positive integer")
    return number


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Everything-backed SQLite FTS5 content index")
    sub = parser.add_subparsers(dest="command", required=True)
    update = sub.add_parser("update", help="incrementally update the content index from Everything HTTP")
    update.add_argument("--db")
    update.add_argument("--endpoint")
    update.add_argument("--query", help="advanced Everything query; incompatible with --prune and --drives")
    update.add_argument("--drives", help="comma-separated drive letters; defaults to config/search-config.json or the system drive")
    update.add_argument("--extensions", default=DEFAULT_EXTENSIONS)
    update.add_argument("--page-size", type=positive_int, default=500)
    update.add_argument("--progress-every", type=positive_int, default=25)
    update.add_argument("--max-mb", type=positive_int, default=100)
    update.add_argument("--prune", action="store_true", help="remove DB entries absent from this full result set")
    update.set_defaults(func=cmd_update)
    search = sub.add_parser("search", help="search previously extracted file contents without network access")
    search.add_argument("query")
    search.add_argument("--db")
    search.add_argument("--extensions")
    search.add_argument("--path")
    search.add_argument("--limit", type=positive_int, default=20)
    search.set_defaults(func=cmd_search)
    status = sub.add_parser("status", help="show index health and coverage")
    status.add_argument("--db")
    status.set_defaults(func=cmd_status)
    return parser


def main() -> int:
    try:
        args = build_parser().parse_args()
        return args.func(args)
    except Exception as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
