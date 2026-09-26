"""Regression checks for standalone distribution and data-integrity failures."""
import contextlib
import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
from unittest.mock import patch

from test_find import index, ROOT


class ReleaseReviewTests(unittest.TestCase):
    def test_installed_skill_includes_license(self):
        license_text = (ROOT / "skills/find/LICENSE").read_text()
        self.assertIn("MIT License", license_text)
        self.assertIn("Copyright (c) 2026 Find contributors", license_text)
        self.assertIn("The above copyright notice and this permission notice", license_text)

    def test_cp949_even_byte_text_and_utf16(self):
        text = "드론안전"
        self.assertEqual(index.decode_text(text.encode("cp949")), text)
        self.assertEqual(index.decode_text(text.encode("utf-16")), text)

    def test_malformed_response_cannot_prune_existing_index(self):
        with tempfile.TemporaryDirectory() as folder:
            database = Path(folder) / "index.db"
            document = Path(folder) / "sample.txt"
            document.write_text("retained content", encoding="utf-8")
            args = index.build_parser().parse_args(["update", "--db", str(database), "--extensions", "txt"])
            with patch.object(index, "everything_results", return_value=iter([(document, 1)])), contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                index.cmd_update(args)
            args.prune = True
            with patch.object(index.Path, "is_dir", return_value=True), patch.object(index, "open_everything_url", return_value=io.BytesIO(b'{}')), contextlib.redirect_stdout(io.StringIO()):
                with self.assertRaises(RuntimeError):
                    index.cmd_update(args)
            with contextlib.closing(index.connect(database, writable=False)) as db:
                self.assertEqual(db.execute("SELECT COUNT(*) FROM files").fetchone()[0], 1)

    def test_content_endpoint_requires_loopback(self):
        for endpoint in ["https://example.invalid/", "file:///C:/private", "http://user:pass@localhost/"]:
            with self.subTest(endpoint=endpoint), self.assertRaises(ValueError):
                index.resolve_endpoint(endpoint)

    def test_absolute_folder_scope_excludes_similarly_named_sibling(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            wanted = base / "course"
            sibling = base / "course-old"
            wanted.mkdir()
            sibling.mkdir()
            files = [wanted / "a.txt", sibling / "b.txt"]
            for file in files:
                file.write_text("sharedterm", encoding="utf-8")
            database = base / "index.db"
            args = index.build_parser().parse_args(["update", "--db", str(database), "--extensions", "txt"])
            with patch.object(index, "everything_results", return_value=iter((file, 2) for file in files)), contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                index.cmd_update(args)
            search = index.build_parser().parse_args(["search", "sharedterm", "--db", str(database), "--path", str(wanted)])
            with contextlib.redirect_stdout(io.StringIO()) as output:
                index.cmd_search(search)
            data = json.loads(output.getvalue())
            self.assertEqual(data["count"], 1)
            self.assertEqual(data["results"][0]["path"], str(files[0]))

    def test_unavailable_drive_blocks_prune_before_database_creation(self):
        with tempfile.TemporaryDirectory() as folder:
            database = Path(folder) / "index.db"
            args = index.build_parser().parse_args(["update", "--db", str(database), "--prune"])
            with patch.object(index.Path, "is_dir", return_value=False), self.assertRaises(ValueError):
                index.cmd_update(args)
            self.assertFalse(database.exists())

    def test_pagination_changes_and_duplicates_are_rejected(self):
        item = {"path": str(ROOT), "name": "sample.txt"}
        first = {"totalResults": 2, "results": [item]}
        for second in [{"totalResults": 0, "results": []}, first]:
            replies = [io.BytesIO(json.dumps(page).encode()) for page in (first, second)]
            with self.subTest(second=second), patch.object(index, "open_everything_url", side_effect=replies):
                with self.assertRaises(RuntimeError):
                    list(index.everything_results("http://127.0.0.1/", "file:", 1))

    def test_http_redirects_are_rejected(self):
        with self.assertRaises(RuntimeError):
            index.NoEverythingRedirect().redirect_request(None, None, 302, "Found", {}, "https://example.invalid/")

    def test_nonpositive_cli_limits_are_rejected(self):
        for arguments in [["update", "--page-size", "0"], ["update", "--progress-every", "0"], ["search", "topic", "--limit", "-1"]]:
            with self.subTest(arguments=arguments), contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                index.build_parser().parse_args(arguments)

    def test_modern_document_extractors(self):
        try:
            from docx import Document
            from pptx import Presentation
            from pptx.util import Inches
            from openpyxl import Workbook
            from pypdf import PdfWriter
            from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        except ImportError as exc:
            self.skipTest(str(exc))
        text = "flight safety lesson"
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            word = Document()
            word.add_paragraph(text)
            word.save(base / "word.docx")
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = text
            deck.save(base / "slides.pptx")
            workbook = Workbook()
            workbook.active["A1"] = text
            workbook.save(base / "sheet.xlsx")
            workbook.close()
            with zipfile.ZipFile(base / "hangul.hwpx", "w") as archive:
                archive.writestr("Contents/section0.xml", f"<section><p>{text}</p></section>")
            pdf = PdfWriter()
            page = pdf.add_blank_page(width=300, height=300)
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 20 20 Td ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = stream
            pdf.write(base / "document.pdf")
            for file in base.iterdir():
                with self.subTest(format=file.suffix):
                    self.assertIn(text, index.extract_content(file, {})[0])


if __name__ == "__main__":
    unittest.main()
